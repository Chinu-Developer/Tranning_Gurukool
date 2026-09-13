import os
import time
import uuid
import openpyxl

from getpass import getpass
from datetime import datetime
from pyfiglet import figlet_format

from rich.console import Console
from rich.table import Table
from rich.progress import track


# ---------------------------------------------------------
# BASIC SETTINGS
# ---------------------------------------------------------

FILE_NAME = "bank_records.xlsx"

console = Console()


# ---------------------------------------------------------
# CREATE EXCEL FILE IF IT DOES NOT EXIST
# ---------------------------------------------------------

def create_excel_file():

    if not os.path.exists(FILE_NAME):

        workbook = openpyxl.Workbook()

        sheet = workbook.active
        sheet.title = "Bank Records"

        sheet.append([
            "Account No",
            "Name",
            "PIN",
            "Transaction ID",
            "Transaction Type",
            "Amount",
            "Previous Balance",
            "Current Balance",
            "Date-Time"
        ])

        workbook.save(FILE_NAME)
        workbook.close()

        console.print("[green]Excel database created successfully![/green]")


# ---------------------------------------------------------
# BANK BANNER
# ---------------------------------------------------------

def show_banner():

    print(figlet_format("PY BANK"))

    console.print(
        "[bold cyan]Welcome to Python Banking System[/bold cyan]"
    )


# ---------------------------------------------------------
# SMALL LOADING EFFECT
# ---------------------------------------------------------

def loading(message):

    console.print(message)

    for _ in track(range(20), description="Processing..."):
        time.sleep(0.03)


# ---------------------------------------------------------
# CHECK WHETHER ACCOUNT ALREADY EXISTS
# ---------------------------------------------------------

def account_exists(account_no):

    workbook = openpyxl.load_workbook(FILE_NAME)
    sheet = workbook.active

    found = False

    for row in sheet.iter_rows(min_row=2, values_only=True):

        if str(row[0]) == str(account_no):

            found = True
            break

    workbook.close()

    return found


# ---------------------------------------------------------
# CREATE ACCOUNT
# ---------------------------------------------------------

def create_account():

    console.print("\n[bold yellow]CREATE ACCOUNT[/bold yellow]")

    name = input("Enter Name: ")

    account_no = input("Enter Account Number: ")

    # Account number validation
    if not account_no.isdigit():

        console.print("[red]Account number must contain only digits.[/red]")
        return

    if account_exists(account_no):

        console.print(
            "[red]This account number is already taken.[/red]"
        )
        return

    pin = getpass("Create 4-digit PIN: ")
    
    if len(pin) != 4 or not pin.isdigit():

        console.print("[red]PIN must be exactly 4 digits.[/red]")
        return

    try:

        opening_balance = float(
            input("Enter Opening Balance: ")
        )

        if opening_balance < 0:

            console.print(
                "[red]Opening balance cannot be negative.[/red]"
            )
            return

    except ValueError:

        console.print("[red]Please enter valid amount.[/red]")
        return

    transaction_id = str(uuid.uuid4())[:8]

    date_time = datetime.now().strftime(
        "%d-%m-%Y %H:%M:%S"
    )

    workbook = openpyxl.load_workbook(FILE_NAME)

    sheet = workbook.active
    
    sheet.append([
        account_no,
        name,
        pin,
        transaction_id,
        "Opening",
        opening_balance,
        0,
        opening_balance,
        date_time
    ])

    workbook.save(FILE_NAME)

    workbook.close()

    loading("Creating your bank account...")

    console.print(
        "\n[bold green]Account Created Successfully![/bold green]"
    )

    console.print(
        f"Account Number: [cyan]{account_no}[/cyan]"
    )


# ---------------------------------------------------------
# LOGIN
# ---------------------------------------------------------

def login():

    console.print("\n[bold yellow]LOGIN[/bold yellow]")

    account_no = input("Enter Account Number: ")

    pin = getpass("Enter PIN: ")

    workbook = openpyxl.load_workbook(FILE_NAME)

    sheet = workbook.active

    user_name = None

    for row in sheet.iter_rows(min_row=2, values_only=True):

        if (
            str(row[0]) == str(account_no)
            and str(row[2]) == str(pin)
        ):

            user_name = row[1]

            break

    workbook.close()

    if user_name:

        loading("Verifying account...")

        console.print(
            f"\n[bold green]Login Successful![/bold green]"
        )

        console.print(
            f"Welcome [cyan]{user_name}[/cyan]"
        )

        return account_no

    else:

        console.print(
            "[bold red]Invalid Account Number or PIN[/bold red]"
        )

        return None


# ---------------------------------------------------------
# GET CURRENT BALANCE
# ---------------------------------------------------------

def get_balance(account_no):

    workbook = openpyxl.load_workbook(FILE_NAME)

    sheet = workbook.active

    current_balance = None

    # Every matching row updates current_balance.
    # Therefore the last matching row becomes
    # the latest balance.

    for row in sheet.iter_rows(min_row=2, values_only=True):

        if str(row[0]) == str(account_no):

            current_balance = row[7]

    workbook.close()

    return current_balance


# ---------------------------------------------------------
# GET USER NAME AND PIN
# ---------------------------------------------------------

def get_user_details(account_no):

    workbook = openpyxl.load_workbook(FILE_NAME)

    sheet = workbook.active

    name = None
    pin = None

    for row in sheet.iter_rows(min_row=2, values_only=True):

        if str(row[0]) == str(account_no):

            name = row[1]
            pin = row[2]

            break


    workbook.close()

    return name, pin


# ---------------------------------------------------------
# CHECK BALANCE
# ---------------------------------------------------------

def check_balance(account_no):

    balance = get_balance(account_no)

    console.print(
        f"\n[bold cyan]Current Balance: ₹{balance:.2f}[/bold cyan]"
    )


# ---------------------------------------------------------
# DEPOSIT MONEY
# ---------------------------------------------------------

def deposit(account_no):

    console.print("\n[bold yellow]DEPOSIT MONEY[/bold yellow]")

    try:

        amount = float(
            input("Enter Deposit Amount: ")
        )

        if amount <= 0:

            console.print(
                "[red]Deposit amount must be greater than 0.[/red]"
            )
            return

    except ValueError:

        console.print("[red]Enter valid amount.[/red]")
        return

    previous_balance = get_balance(account_no)

    new_balance = previous_balance + amount

    name, pin = get_user_details(account_no)
    if name is None:
        console.print("[red]Seems Data is corrupted name and pin doesn't exists but account does.[/red]")
        return

    transaction_id = str(uuid.uuid4())[:8]

    date_time = datetime.now().strftime(
        "%d-%m-%Y %H:%M:%S"
    )

    workbook = openpyxl.load_workbook(FILE_NAME)

    sheet = workbook.active

    sheet.append([
        account_no,
        name,
        pin,
        transaction_id,
        "Deposit",
        amount,
        previous_balance,
        new_balance,
        date_time
    ])

    workbook.save(FILE_NAME)

    workbook.close()

    loading("Depositing money...")

    console.print(
        f"[green]₹{amount:.2f} deposited successfully.[/green]"
    )

    console.print(
        f"New Balance: [bold cyan]₹{new_balance:.2f}[/bold cyan]"
    )


# ---------------------------------------------------------
# WITHDRAW MONEY
# ---------------------------------------------------------

def withdraw(account_no):

    console.print("\n[bold yellow]WITHDRAW MONEY[/bold yellow]")

    previous_balance = get_balance(account_no)

    try:

        amount = float(
            input("Enter Withdrawal Amount: ")
        )

        if amount <= 0:

            console.print(
                "[red]Withdrawal amount must be greater than 0.[/red]"
            )
            return

    except ValueError:

        console.print("[red]Enter valid amount.[/red]")
        return

    if amount > previous_balance:

        console.print(
            "[bold red]Insufficient Balance![/bold red]"
        )

        return

    new_balance = previous_balance - amount

    name, pin = get_user_details(account_no)

    transaction_id = str(uuid.uuid4())[:8]

    date_time = datetime.now().strftime(
        "%d-%m-%Y %H:%M:%S"
    )

    workbook = openpyxl.load_workbook(FILE_NAME)

    sheet = workbook.active

    sheet.append([
        account_no,
        name,
        pin,
        transaction_id,
        "Withdraw",
        amount,
        previous_balance,
        new_balance,
        date_time
    ])

    workbook.save(FILE_NAME)

    workbook.close()

    loading("Processing withdrawal...")

    console.print(
        f"[green]₹{amount:.2f} withdrawn successfully.[/green]"
    )

    console.print(
        f"Remaining Balance: [bold cyan]₹{new_balance:.2f}[/bold cyan]"
    )


# ---------------------------------------------------------
# TRANSACTION HISTORY
# ---------------------------------------------------------

def transaction_history(account_no):

    workbook = openpyxl.load_workbook(FILE_NAME)

    sheet = workbook.active

    table = Table(
        title="Transaction History"
    )

    table.add_column("Transaction ID")
    table.add_column("Type")
    table.add_column("Amount")
    table.add_column("Previous")
    table.add_column("Balance")
    table.add_column("Date")

    found = False

    for row in sheet.iter_rows(min_row=2, values_only=True):

        if str(row[0]) == str(account_no):

            found = True

            table.add_row(
                str(row[3]),
                str(row[4]),
                f"₹{row[5]:.2f}",
                f"₹{row[6]:.2f}",
                f"₹{row[7]:.2f}",
                str(row[8])
            )

    workbook.close()

    if found:

        console.print(table)

    else:

        console.print(
            "[red]No transactions found.[/red]"
        )


# ---------------------------------------------------------
# USER MENU AFTER LOGIN
# ---------------------------------------------------------

def banking_menu(account_no):
    #
    while True:

        console.print(f"\n[bold cyan]BANKING OPTIONS for {account_no} [/bold cyan]")

        print("1. Check Balance")
        print("2. Deposit Money")
        print("3. Withdraw Money")
        print("4. Transaction History")
        print("5. Logout")

        choice = input("\nEnter Choice: ")

        if choice == "1":

            check_balance(account_no)

        elif choice == "2":

            deposit(account_no)

        elif choice == "3":

            withdraw(account_no)

        elif choice == "4":

            transaction_history(account_no)

        elif choice == "5":

            console.print(
                "\n[yellow]Logged out successfully.[/yellow]"
            )

            break

        else:

            console.print(
                "[red]Invalid Choice[/red]"
            )


# ---------------------------------------------------------
# MAIN PROGRAM
# ---------------------------------------------------------

create_excel_file()

while True:

    show_banner()

    print("1. Create Account")
    print("2. Login")
    print("3. Exit")

    choice = input("\nEnter Choice: ")

    if choice == "1":

        create_account()

    elif choice == "2":

        logged_in_account = login()

        if logged_in_account:

            banking_menu(logged_in_account)

    elif choice == "3":

        console.print(
            "\n[bold cyan]Thank you for using PY BANK![/bold cyan]"
        )

        break

    else:

        console.print(
            "[red]Invalid Choice[/red]"
        )